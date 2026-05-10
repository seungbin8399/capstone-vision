import csv
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path.cwd()
IMG_DIR = next(
    p for p in ROOT.iterdir()
    if p.is_dir() and any(ord(ch) > 127 for ch in p.name)
)
IMAGES = sorted(
    [p for p in IMG_DIR.iterdir() if p.suffix.lower() in [".png", ".jpg", ".jpeg"]],
    key=lambda p: p.name,
)

OUT = ROOT / "RealSense_유압블록_비전실험_중간발표.pptx"
CSV_PATH = ROOT / "realsense_depth_log_without_postit.csv"
PLOT_PATH = ROOT / "depth_stability_plot.png"
DIAGRAM_PATH = ROOT / "camera_xyz_diagram.png"


def load_summary():
    with CSV_PATH.open("r", encoding="utf-8") as file:
        rows = list(csv.DictReader(file))

    def values(key):
        result = []
        for row in rows:
            try:
                result.append(float(row[key]))
            except Exception:
                pass
        return result

    raw = values("raw_depth_m")
    filtered = values("filtered_depth_m")
    camera_x = values("camera_x_m")
    camera_y = values("camera_y_m")
    camera_z = values("camera_z_m")

    return rows, {
        "n": len(rows),
        "raw_min": min(raw) if raw else 0.412,
        "raw_max": max(raw) if raw else 0.414,
        "filtered_min": min(filtered) if filtered else 0.412,
        "filtered_max": max(filtered) if filtered else 0.414,
        "filtered_avg": sum(filtered) / len(filtered) if filtered else 0.412,
        "camera_x_avg": sum(camera_x) / len(camera_x) if camera_x else -0.005,
        "camera_y_avg": sum(camera_y) / len(camera_y) if camera_y else 0.062,
        "camera_z_avg": sum(camera_z) / len(camera_z) if camera_z else 0.412,
        "center_u": int(float(rows[0]["center_u"])) if rows else 325,
        "center_v": int(float(rows[0]["center_v"])) if rows else 335,
        "sample_u": int(float(rows[0]["sample_u"])) if rows else 305,
        "sample_v": int(float(rows[0]["sample_v"])) if rows else 315,
        "raw": raw,
        "filtered": filtered,
    }


ROWS, SUMMARY = load_summary()


def make_depth_plot():
    width, height = 1100, 520
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    margin_l, margin_r, margin_t, margin_b = 90, 40, 55, 75
    plot_w = width - margin_l - margin_r
    plot_h = height - margin_t - margin_b
    raw = SUMMARY["raw"]
    filtered = SUMMARY["filtered"]

    draw.rectangle([0, 0, width - 1, height - 1], outline=(220, 226, 235), width=2)
    draw.text((40, 18), "Depth Stability Log (without_postit)", fill=(25, 40, 65))
    draw.text((width - 300, 18), f"n={SUMMARY['n']}, rolling median window=5", fill=(80, 90, 105))

    if not raw:
        draw.text((width // 2 - 90, height // 2), "No CSV data", fill=(160, 0, 0))
        image.save(PLOT_PATH)
        return

    y_min = min(min(raw), min(filtered)) - 0.001
    y_max = max(max(raw), max(filtered)) + 0.001
    if y_max <= y_min:
        y_max = y_min + 0.01

    for index in range(6):
        y = margin_t + plot_h * index / 5
        value = y_max - (y_max - y_min) * index / 5
        draw.line([(margin_l, y), (width - margin_r, y)], fill=(232, 236, 242), width=1)
        draw.text((20, y - 8), f"{value:.3f}m", fill=(80, 90, 105))

    draw.line([(margin_l, margin_t), (margin_l, height - margin_b)], fill=(70, 80, 95), width=2)
    draw.line([(margin_l, height - margin_b), (width - margin_r, height - margin_b)], fill=(70, 80, 95), width=2)

    def point(i, value):
        x = margin_l + plot_w * i / max(1, len(raw) - 1)
        y = margin_t + plot_h * (y_max - value) / (y_max - y_min)
        return x, y

    for series, color in [(raw, (230, 105, 70)), (filtered, (35, 130, 210))]:
        points = [point(i, value) for i, value in enumerate(series)]
        for a, b in zip(points, points[1:]):
            draw.line([a, b], fill=color, width=3)
        for x, y in points:
            draw.ellipse([x - 3, y - 3, x + 3, y + 3], fill=color)

    draw.rectangle([margin_l + 20, margin_t + 20, margin_l + 270, margin_t + 86], fill=(255, 255, 255), outline=(210, 215, 225))
    draw.line([(margin_l + 38, margin_t + 42), (margin_l + 95, margin_t + 42)], fill=(230, 105, 70), width=4)
    draw.text((margin_l + 110, margin_t + 32), "raw_depth_m", fill=(45, 55, 70))
    draw.line([(margin_l + 38, margin_t + 68), (margin_l + 95, margin_t + 68)], fill=(35, 130, 210), width=4)
    draw.text((margin_l + 110, margin_t + 58), "filtered_depth_m", fill=(45, 55, 70))
    draw.text((margin_l, height - 48), "sample index", fill=(80, 90, 105))
    draw.text((width - 365, height - 48), f"filtered range: {SUMMARY['filtered_min']:.3f}~{SUMMARY['filtered_max']:.3f}m", fill=(25, 40, 65))

    image.save(PLOT_PATH)


def make_camera_diagram():
    width, height = 1100, 520
    image = Image.new("RGB", (width, height), (250, 252, 255))
    draw = ImageDraw.Draw(image)

    draw.rectangle([0, 0, width - 1, height - 1], outline=(220, 226, 235), width=2)
    draw.text((40, 25), "Pixel + Depth -> Camera Frame 3D Point", fill=(25, 40, 65))

    draw.rounded_rectangle([70, 155, 260, 330], radius=18, fill=(229, 239, 255), outline=(45, 95, 165), width=3)
    draw.ellipse([135, 195, 195, 255], fill=(45, 95, 165))
    draw.text((105, 345), "RealSense D435", fill=(25, 40, 65))

    draw.rectangle([415, 110, 730, 365], fill=(255, 255, 255), outline=(80, 100, 130), width=3)
    draw.text((500, 382), "color image plane", fill=(80, 90, 105))

    draw.ellipse([564, 252, 578, 266], fill=(220, 35, 35))
    draw.text((585, 248), f"center=({SUMMARY['center_u']},{SUMMARY['center_v']})", fill=(150, 30, 30))
    draw.rectangle([520, 205, 560, 245], outline=(0, 165, 80), width=3)
    draw.ellipse([537, 222, 547, 232], fill=(0, 165, 80))
    draw.text((395, 185), f"sample=({SUMMARY['sample_u']},{SUMMARY['sample_v']})", fill=(0, 120, 75))

    draw.line([(270, 240), (405, 240)], fill=(70, 80, 95), width=4)
    draw.polygon([(405, 240), (385, 230), (385, 250)], fill=(70, 80, 95))
    draw.line([(735, 240), (870, 240)], fill=(70, 80, 95), width=4)
    draw.polygon([(870, 240), (850, 230), (850, 250)], fill=(70, 80, 95))

    draw.rounded_rectangle([880, 160, 1050, 320], radius=16, fill=(235, 248, 240), outline=(50, 135, 85), width=3)
    draw.text((905, 190), "camera_xyz", fill=(25, 80, 55))
    draw.text((905, 225), f"X={SUMMARY['camera_x_avg']:.3f}m", fill=(25, 80, 55))
    draw.text((905, 255), f"Y={SUMMARY['camera_y_avg']:.3f}m", fill=(25, 80, 55))
    draw.text((905, 285), f"Z={SUMMARY['camera_z_avg']:.3f}m", fill=(25, 80, 55))

    draw.text((390, 430), "X=(u-cx)*Z/fx    Y=(v-cy)*Z/fy    Z=filtered_depth_m", fill=(25, 40, 65))
    image.save(DIAGRAM_PATH)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.25), Inches(0.6))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = "Malgun Gothic"
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(22, 36, 60)

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(0.9), Inches(12), Inches(0.32))
        paragraph = sub.text_frame.paragraphs[0]
        paragraph.text = subtitle
        paragraph.font.name = "Malgun Gothic"
        paragraph.font.size = Pt(12)
        paragraph.font.color.rgb = RGBColor(92, 106, 125)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.22), Inches(12.25), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(219, 226, 235)
    line.line.fill.background()


def add_bullets(slide, bullets, x, y, w, h, font_size=16):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    text_frame = box.text_frame
    text_frame.clear()
    text_frame.word_wrap = True

    for index, bullet in enumerate(bullets):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = bullet
        paragraph.font.name = "Malgun Gothic"
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor(35, 45, 60)
        paragraph.space_after = Pt(8)


def add_script(slide, script):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(6.22), Inches(12.25), Inches(0.9))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(239, 244, 250)
    shape.line.color.rgb = RGBColor(219, 226, 235)

    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.margin_left = Inches(0.18)
    text_frame.margin_right = Inches(0.18)
    text_frame.margin_top = Inches(0.08)

    paragraph = text_frame.paragraphs[0]
    paragraph.text = "발표 대본: " + script
    paragraph.font.name = "Malgun Gothic"
    paragraph.font.size = Pt(10.5)
    paragraph.font.color.rgb = RGBColor(45, 58, 78)


def add_picture_fit(slide, path, x, y, w, h):
    image = Image.open(path)
    image_w, image_h = image.size
    box_ratio = w / h
    image_ratio = image_w / image_h

    if image_ratio >= box_ratio:
        pic_w = w
        pic_h = w / image_ratio
    else:
        pic_h = h
        pic_w = h * image_ratio

    pic_x = x + (w - pic_w) / 2
    pic_y = y + (h - pic_h) / 2

    return slide.shapes.add_picture(str(path), Inches(pic_x), Inches(pic_y), width=Inches(pic_w), height=Inches(pic_h))


def add_label(slide, text, x, y, w, h, fill, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    shape.line.color.rgb = RGBColor(219, 226, 235)
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.margin_left = Inches(0.12)
    text_frame.margin_right = Inches(0.12)
    text_frame.margin_top = Inches(0.05)
    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Malgun Gothic"
    paragraph.font.size = Pt(12)
    paragraph.font.bold = True
    paragraph.font.color.rgb = color
    paragraph.alignment = PP_ALIGN.CENTER


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(248, 250, 252)
    return slide


def build_presentation():
    make_depth_plot()
    make_camera_diagram()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    accent = RGBColor(24, 96, 170)
    green = RGBColor(36, 130, 80)

    # Slide 1
    slide = blank_slide(prs)
    add_title(slide, "RealSense D435 기반 유압블록 비전 실험", "캡스톤 중간 진행 발표 | AI 기반 비정형 유압블록 로봇팔 자동 적재 시스템")
    add_bullets(slide, [
        "목표: 유압블록 중심 후보 pixel과 안정화된 depth 추출",
        "현재 단계: Windows + Python + OpenCV + pyrealsense2 기반 실험",
        "출력: center pixel, depth, camera frame 3D 좌표, CSV 로그",
        "향후 ROS2 PoseStamped 및 robot base_link 좌표 변환으로 확장",
    ], 0.75, 1.55, 5.4, 3.1, 16)
    add_picture_fit(slide, IMAGES[0], 6.45, 1.45, 5.9, 3.7)
    add_label(slide, "실험 장비: Intel RealSense D435 + 실제 유압블록", 6.85, 5.25, 5.1, 0.42, (235, 248, 255), accent)
    add_script(slide, "이번 발표에서는 전체 자동 적재 시스템 중 비전 파트의 현재 진행상황을 설명드리겠습니다. 현재는 RealSense D435로 유압블록의 중심 후보와 depth를 안정적으로 얻고, camera frame 3D 좌표까지 계산했습니다.")

    # Slide 2
    slide = blank_slide(prs)
    add_title(slide, "실험 환경 및 조건")
    add_bullets(slide, [
        "환경: Windows 데스크탑, Intel RealSense D435, 실제 금속 유압블록",
        "라이브러리: pyrealsense2, OpenCV, NumPy",
        "RGB/Depth stream 정상 수신 및 aligned depth 사용",
        "현재 조건: 포스트잇 없이 금속 유압블록만 사용",
        "실험 데이터는 CSV로 저장하여 raw/filtered depth 비교",
    ], 0.75, 1.5, 5.6, 3.8, 15)
    add_picture_fit(slide, IMAGES[1], 6.5, 1.4, 5.7, 3.2)
    add_picture_fit(slide, IMAGES[8], 6.5, 4.25, 5.7, 1.55)
    add_script(slide, "금속 표면은 반사가 심하고 구멍이 많아서 depth가 깨지는 구간이 있었습니다. 그래서 초기에는 자동 인식보다 안정적인 depth 측정과 로그 확보에 초점을 맞췄습니다.")

    # Slide 3
    slide = blank_slide(prs)
    add_title(slide, "초기 시행착오: 자동 contour/depth mask의 한계")
    add_bullets(slide, [
        "Depth mask 방식: 금속 유압블록 depth가 중간중간 깨짐",
        "배경/책상 depth가 더 안정적으로 잡혀 큰 contour가 생성됨",
        "Color contour 방식: 반사와 구멍 때문에 mask가 조각남",
        "결론: 현재 단계에서는 fixed ROI 기반 측정이 더 안정적",
    ], 0.75, 1.55, 5.15, 3.6, 15)
    add_picture_fit(slide, IMAGES[4], 6.15, 1.35, 6.15, 2.15)
    add_picture_fit(slide, IMAGES[5], 6.15, 3.8, 6.15, 1.75)
    add_label(slide, "검출 자동화보다 안정적인 측정값 확보를 우선", 0.95, 5.22, 4.65, 0.5, (255, 245, 230), RGBColor(150, 80, 20))
    add_script(slide, "처음에는 depth mask와 color contour를 적용했지만, 금속 표면의 특성 때문에 bbox가 크게 잡히거나 중심점이 흔들렸습니다. 그래서 중간 단계에서는 fixed ROI로 안정적인 측정 기준을 먼저 만들었습니다.")

    # Slide 4
    slide = blank_slide(prs)
    add_title(slide, "현재 적용한 fixed ROI 방식")
    add_bullets(slide, [
        "파란 박스: manual ROI, 유압블록 전체 영역",
        "빨간 점: center_u, center_v = 유압블록 중심 후보 pixel",
        "초록 박스/점: sample_u, sample_v = 실제 depth 측정 위치",
        "중심 구멍과 반사를 피하기 위해 중심 후보와 depth sampling 위치를 분리",
        "최근 5개 valid depth의 rolling median으로 filtered_depth 계산",
    ], 0.75, 1.45, 5.5, 3.9, 14.5)
    add_picture_fit(slide, IMAGES[8], 6.45, 1.35, 5.8, 3.3)
    add_label(slide, "빨간 점은 물체 중심 후보, 초록 점은 depth 읽는 지점", 6.78, 4.95, 5.1, 0.5, (235, 248, 240), green)
    add_script(slide, "현재 방식에서는 중심 후보와 depth를 읽는 위치를 의도적으로 분리했습니다. 로봇에는 중심 후보를 넘기되, depth는 구멍이나 반사가 적은 표면에서 읽어 안정성을 높였습니다.")

    # Slide 5
    slide = blank_slide(prs)
    add_title(slide, "안정 설정 및 로그 결과")
    add_bullets(slide, [
        "Manual ROI: X1=230, Y1=230, X2=420, Y2=440",
        "Depth sampling offset: X=-20, Y=-20",
        "Depth sample radius: 10 px, 즉 21x21 영역 median",
        f"center pixel: ({SUMMARY['center_u']}, {SUMMARY['center_v']}) / sample pixel: ({SUMMARY['sample_u']}, {SUMMARY['sample_v']})",
        f"raw_depth: {SUMMARY['raw_min']:.3f}~{SUMMARY['raw_max']:.3f}m",
        f"filtered_depth: {SUMMARY['filtered_min']:.3f}~{SUMMARY['filtered_max']:.3f}m",
    ], 0.72, 1.42, 5.25, 4.4, 14.2)
    add_picture_fit(slide, PLOT_PATH, 6.2, 1.35, 6.1, 3.55)
    add_label(slide, "depth 흔들림 약 2~3mm 수준으로 안정화", 6.85, 5.15, 4.8, 0.48, (235, 248, 255), accent)
    add_script(slide, "현재 without_postit 조건에서 raw depth와 filtered depth가 0.412에서 0.414m 부근으로 안정적으로 측정되었습니다. invalid depth도 거의 없어서 다음 좌표 변환 단계로 넘어갈 수 있는 상태입니다.")

    # Slide 6
    slide = blank_slide(prs)
    add_title(slide, "Camera Frame 3D 좌표 변환")
    add_bullets(slide, [
        "RealSense color stream profile에서 intrinsics를 직접 읽어 사용",
        "fx=604.475, fy=603.970, cx=332.253, cy=244.044",
        "공식: X=(u-cx)Z/fx, Y=(v-cy)Z/fy, Z=filtered_depth",
        f"결과 평균: X={SUMMARY['camera_x_avg']:.3f}m, Y={SUMMARY['camera_y_avg']:.3f}m, Z={SUMMARY['camera_z_avg']:.3f}m",
        "주의: 현재 좌표는 robot base_link가 아니라 camera frame 기준",
    ], 0.75, 1.45, 5.5, 4.1, 14.5)
    add_picture_fit(slide, DIAGRAM_PATH, 6.35, 1.45, 5.85, 3.2)
    add_label(slide, f"camera_xyz≈({SUMMARY['camera_x_avg']:.3f}, {SUMMARY['camera_y_avg']:.3f}, {SUMMARY['camera_z_avg']:.3f})m", 6.95, 5.05, 4.65, 0.5, (235, 248, 240), green)
    add_script(slide, "pixel 좌표와 filtered depth를 이용해 카메라 기준 3D 좌표를 계산했습니다. 현재 결과는 카메라 기준 위치이며, 실제 로봇 제어에는 base_link 기준 좌표로 변환하는 과정이 추가로 필요합니다.")

    # Slide 7
    slide = blank_slide(prs)
    add_title(slide, "현재 한계와 다음 단계")
    add_bullets(slide, [
        "현재 한계",
        "manual ROI 기반이므로 완전 자동 검출은 아님",
        "금속 표면 반사와 구멍 때문에 샘플링 위치에 민감함",
        "현재 좌표는 camera frame 기준이며 robot base_link 기준이 아님",
        "다음 단계",
        "ROS2 비전 노드로 포팅하고 PoseStamped publish",
        "camera_xyz → robot base_link 좌표 변환(TF/extrinsic calibration)",
        "연구실 환경에서 조명, 카메라 각도, 거리 재측정",
        "OMY-F3M 로봇팔 motion planning 목표점과 연동",
    ], 0.75, 1.35, 6.1, 4.65, 13.7)
    pipeline_steps = [
        ("RealSense\nRGB/Depth", 7.15, 1.55),
        ("Vision Node\nfixed ROI", 9.7, 1.55),
        ("PoseStamped\n/detection", 7.15, 2.8),
        ("TF 변환\nbase_link", 9.7, 2.8),
        ("OMY-F3M\nMotion", 7.15, 4.05),
    ]
    for index, (text, x, y) in enumerate(pipeline_steps):
        add_label(slide, text, x, y, 2.1, 0.78, (235, 248, 255) if index < 3 else (235, 248, 240), accent if index < 3 else green)
    add_script(slide, "현재는 비전 측정값을 안정화한 단계입니다. 다음에는 ROS2에서 PoseStamped 형태로 publish하고, 카메라 좌표계를 로봇 base_link 좌표계로 변환한 뒤 OMY-F3M 로봇팔의 목표 위치로 연결하겠습니다.")

    prs.save(OUT)


if __name__ == "__main__":
    build_presentation()
    print(OUT)
    print(PLOT_PATH)
    print(DIAGRAM_PATH)
